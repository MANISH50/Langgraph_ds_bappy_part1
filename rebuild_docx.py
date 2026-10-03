from pathlib import Path
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt

md_path = Path('LANGGRAPH_STUDY_GUIDE.md')
docx_path = Path('LANGGRAPH_STUDY_GUIDE.docx')
text = md_path.read_text(encoding='utf-8')
lines = text.splitlines()

doc = Document()

in_code = False
code_lines = []

for raw in lines:
    line = raw.rstrip()
    stripped = line.strip()

    if line.startswith('```'):
        if in_code:
            p = doc.add_paragraph()
            p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            p.paragraph_format.left_indent = 0
            run = p.add_run('\n'.join(code_lines))
            run.font.name = 'Consolas'
            run.font.size = Pt(16)
            run.font.color.rgb = None
            in_code = False
            code_lines = []
        else:
            in_code = True
            code_lines = []
        continue

    if in_code:
        code_lines.append(line)
        continue

    if not stripped:
        doc.add_paragraph('')
        continue

    if stripped.startswith('#'):
        level = len(stripped) - len(stripped.lstrip('#'))
        heading = stripped[level:].strip()
        h = doc.add_heading(heading, level=min(level, 3))
        h.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        continue

    if stripped.startswith('- ') or stripped.startswith('* '):
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(stripped[2:].strip())
        continue

    if stripped.startswith('> '):
        p = doc.add_paragraph()
        r = p.add_run(stripped[2:].strip())
        r.italic = True
        continue

    doc.add_paragraph(line)

doc.save(docx_path)

doc2 = Document(docx_path)
combined = '\n'.join(p.text for p in doc2.paragraphs)
print('paragraphs=', len(doc2.paragraphs))
print('has_tkinter=', 'import tkinter as tk' in combined)
print('has_annotated=', 'It adds extra instructions to the type.' in combined)
print('has_persistence=', 'persist state, isolate thread state' in combined)
print('code_block_count=', combined.count('import tkinter as tk'))
print('size=', docx_path.stat().st_size)
